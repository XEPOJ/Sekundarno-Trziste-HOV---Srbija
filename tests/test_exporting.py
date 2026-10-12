from datetime import datetime

from openpyxl import load_workbook

from securities_analyzer.api import exporting


def test_unique_export_path_appends_counter_on_collision(tmp_path, monkeypatch):
    class FixedNow(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 10, 12, 9, 30, 15)

    monkeypatch.setattr(exporting, "datetime", FixedNow)
    d = tmp_path / "out"
    p1 = exporting.unique_export_path(d, "securities", "csv")
    p1.write_text("x")
    p2 = exporting.unique_export_path(d, "securities", "csv")
    p2.write_text("x")
    p3 = exporting.unique_export_path(d, "securities", "csv")
    assert [p.name for p in (p1, p2, p3)] == [
        "securities_20261012_093015.csv", "securities_20261012_093015_2.csv",
        "securities_20261012_093015_3.csv"]


def test_write_table_csv_has_bom_header_and_rows(tmp_path):
    path = tmp_path / "t.csv"
    exporting.write_table(path, ["A", "B"], [[1, "x"], [2, None]], "csv")
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    assert path.read_text(encoding="utf-8-sig").splitlines() == ["A,B", "1,x", "2,"]


def test_write_table_xlsx_sheet_header_and_rows(tmp_path):
    path = tmp_path / "t.xlsx"
    exporting.write_table(path, ["A", "B"], [[1, "x"], [2, None]], "xlsx")
    ws = load_workbook(path).active
    assert ws.title == "Podaci"
    assert [[c.value for c in row] for row in ws.iter_rows()] == [["A", "B"], [1, "x"], [2, None]]
