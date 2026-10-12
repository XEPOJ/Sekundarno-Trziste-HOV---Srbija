import time
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from securities_analyzer.api.app import create_app
from securities_analyzer.db.schema import connect
from securities_analyzer.query import service
from tests.factories import load_rows, make_row, rows_to_csv


@pytest.fixture
def client(db_path, tmp_path):
    c = connect(db_path)
    load_rows(c, [
        make_row(isin="RSMFRSD0000A", currency="RSD", turnover=100.0, avg_yield_pct=4.0),
        make_row(isin="RSMFRSD0000B", currency="RSD", turnover=300.0, avg_yield_pct=6.0),
        make_row(isin="RSMFRSD0000C", currency="EUR", turnover=50.0, avg_yield_pct=3.0),
    ])
    c.close()
    # Two rows per year: a single row would trip the "<50% of stored rows" guard for 2023.
    app = create_app(db_path, fetch=lambda y: rows_to_csv(
        [make_row(trade_date=date(y, 6, 10), settlement_date=date(y, 6, 10), isin=i)
         for i in ("RSMFRSD0000Y", "RSMFRSD0000Z")]),
        export_dir=tmp_path / "exports")
    return TestClient(app, base_url="http://127.0.0.1")


def test_meta(client):
    j = client.get("/api/meta").json()
    assert j["rows"] == 3 and j["first_trade_date"] == "2023-06-15" and j["version"]


def test_securities(client):
    assert {s["isin"] for s in client.get("/api/securities").json()} == {
        "RSMFRSD0000A", "RSMFRSD0000B", "RSMFRSD0000C"}


def test_rows_with_filter_and_paging(client):
    body = {"filter": {"currency": ["RSD"]}, "sort": "turnover", "order": "desc", "page": 1, "page_size": 1}
    j = client.post("/api/query/rows", json=body).json()
    assert j["total"] == 2 and j["rows"][0]["turnover"] == 300.0 and j["page"] == 1


def test_rows_default_body_returns_everything(client):
    assert client.post("/api/query/rows", json={}).json()["total"] == 3


def test_invalid_filter_is_422(client):
    r = client.post("/api/query/rows", json={"filter": {"currency": ["rsd"]}})
    assert r.status_code == 422
    r = client.post("/api/query/rows", json={"filter": {"nope": 1}})
    assert r.status_code == 422


def test_bad_sort_is_422(client):
    assert client.post("/api/query/rows", json={"sort": "id; DROP"}).status_code == 422
    assert client.post("/api/query/extremes", json={"sort_by": "isin"}).status_code == 422


def test_summary_is_split_by_currency(client):
    j = client.post("/api/query/summary", json={}).json()
    assert {k["currency"] for k in j["kpis"]} == {"RSD", "EUR"} and j["quality_flagged"] == 0


def test_timeseries_scatter_buckets_extremes(client):
    assert client.post("/api/query/timeseries", json={"granularity": "month"}).json()
    assert len(client.post("/api/query/scatter", json={}).json()) == 3
    assert client.post("/api/query/buckets", json={}).json()
    ex = client.post("/api/query/extremes", json={"filter": {"currency": ["RSD"]}}).json()
    assert ex[0]["highest_pct"] == 6.0 and ex[0]["lowest_pct"] == 4.0
    assert client.post("/api/query/timeseries", json={"granularity": "year"}).status_code == 422


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_export(client, tmp_path, fmt):
    j = client.post("/api/export", json={"format": fmt, "filter": {"currency": ["RSD"]}}).json()
    assert j["rows"] == 2
    p = Path(j["path"])
    assert p.exists() and p.suffix == f".{fmt}" and p.parent == tmp_path / "exports"
    if fmt == "csv":
        assert p.read_bytes().startswith(b"\xef\xbb\xbf")
        lines = p.read_text(encoding="utf-8-sig").splitlines()
        assert lines[0].split(",") == list(service.ROW_COLUMNS) and len(lines) == 3
    else:
        from openpyxl import load_workbook
        wb = load_workbook(p)
        ws = wb["Podaci"]
        assert wb.sheetnames == ["Podaci"]
        data = list(ws.iter_rows(values_only=True))
        assert list(data[0]) == list(service.ROW_COLUMNS) and len(data) == 3
        wb.close()


def test_saved_queries_crud(client):
    f = {"currency": ["EUR"], "yield_pct": {"min": 3.5, "min_exclusive": True}}
    assert client.put("/api/saved-queries", json={"name": "Moj upit", "filter": f}).status_code == 200
    items = client.get("/api/saved-queries").json()
    assert items[0]["name"] == "Moj upit" and items[0]["filter"]["currency"] == ["EUR"]
    assert client.delete("/api/saved-queries/Moj upit").status_code == 200
    assert client.get("/api/saved-queries").json() == []
    assert client.put("/api/saved-queries", json={"name": "x", "filter": {"bad": 1}}).status_code == 422


def test_saved_query_name_with_slash_can_be_deleted(client):
    assert client.put("/api/saved-queries", json={"name": "a/b", "filter": {}}).status_code == 200
    assert client.delete("/api/saved-queries/a%2Fb").status_code == 200
    assert client.get("/api/saved-queries").json() == []
    assert client.put("/api/saved-queries", json={"name": "a/b", "filter": {}}).status_code == 200
    assert client.delete("/api/saved-queries/a/b").status_code == 200
    assert client.get("/api/saved-queries").json() == []


def test_update_flow(client):
    assert client.get("/api/update/status").json()["state"] == "idle"
    assert client.post("/api/update").json() == {"started": True}
    for _ in range(100):
        st = client.get("/api/update/status").json()
        if st["state"] != "running":
            break
        time.sleep(0.05)
    assert st["state"] == "done", st
    assert st["report"]["ok"] is True
    assert any("Preuzimam" in m for m in st["messages"])
    assert client.get("/api/meta").json()["rows"] > 3
    assert client.get("/api/meta").json()["last_update_utc"]


def test_meta_has_last_attempt_key(client):
    j = client.get("/api/meta").json()
    assert "last_attempt_utc" in j and j["last_attempt_utc"] is None


def test_untrusted_host_is_rejected(client):
    assert client.get("/api/meta", headers={"Host": "evil.example"}).status_code == 400
    assert client.get("/api/meta", headers={"Host": "localhost:8123"}).status_code == 200


def test_heartbeat_records_time(client):
    assert client.app.state.last_heartbeat is None
    assert client.post("/api/heartbeat").json() == {"ok": True}
    assert isinstance(client.app.state.last_heartbeat, float)


def test_page_upper_bound_is_422(client):
    assert client.post("/api/query/rows", json={"page": 1_000_001}).status_code == 422
    assert client.post("/api/query/rows", json={"page": 1_000_000}).status_code == 200


def test_filter_list_lengths_are_bounded(client):
    ok = {"filter": {"currency": ["RSD"] * 50, "isin": ["RSMFRSD0000A"] * 500}}
    assert client.post("/api/query/rows", json=ok).status_code == 200
    assert client.post("/api/query/rows", json={"filter": {"currency": ["RSD"] * 51}}).status_code == 422
    assert client.post("/api/query/rows", json={"filter": {"isin": ["RSMFRSD0000A"] * 501}}).status_code == 422


@pytest.mark.parametrize("name", ["", "   ", ".", " .. ", "a\x00b", "a\nb", "\t"])
def test_saved_query_bad_names_are_422(client, name):
    assert client.put("/api/saved-queries", json={"name": name, "filter": {}}).status_code == 422


def test_saved_query_name_is_stripped(client):
    assert client.put("/api/saved-queries", json={"name": "  Moj  ", "filter": {}}).json() == {"saved": "Moj"}
    assert client.get("/api/saved-queries").json()[0]["name"] == "Moj"


def test_saved_query_with_pivot_roundtrip(client):
    pivot = {"dimensions": ["isin"], "measures": ["turnover", "quantity"], "sort_by": "turnover",
             "order": "desc", "limit": 10, "chart": {"type": "bar", "horizontal": True}}
    body = {"name": "Top hartije", "filter": {"currency": ["RSD"]}, "pivot": pivot}
    assert client.put("/api/saved-queries", json=body).status_code == 200
    (item,) = client.get("/api/saved-queries").json()
    assert item["name"] == "Top hartije" and item["filter"]["currency"] == ["RSD"]
    assert item["pivot"] == pivot


def test_saved_query_without_pivot_has_null_and_overwrite_clears(client):
    pivot = {"dimensions": [], "measures": ["turnover"], "sort_by": "group", "order": "asc",
             "limit": 100, "chart": {"type": "table", "horizontal": False}}
    client.put("/api/saved-queries", json={"name": "x", "filter": {}, "pivot": pivot})
    assert client.get("/api/saved-queries").json()[0]["pivot"] == pivot
    client.put("/api/saved-queries", json={"name": "x", "filter": {}})
    assert client.get("/api/saved-queries").json()[0]["pivot"] is None


@pytest.mark.parametrize("pivot", [
    {"dimensions": ["nope"]},
    {"measures": []},
    {"sort_by": "quantity"},
    {"chart": {"type": "pie"}},
    {"limit": 99999},
])
def test_saved_query_rejects_invalid_pivot(client, pivot):
    r = client.put("/api/saved-queries", json={"name": "bad", "filter": {}, "pivot": pivot})
    assert r.status_code == 422


def test_pivot_catalog(client):
    j = client.get("/api/pivot/catalog").json()
    assert [d["id"] for d in j["dimensions"]][:2] == ["isin", "date_day"]
    assert "weighted_yield" in [m["id"] for m in j["measures"]]
    assert j["limits"]["max_limit"] == 5000


def test_pivot_query_default_body_totals_per_currency(client):
    j = client.post("/api/query/pivot", json={}).json()
    assert [(r["currency"], r["values"][0]) for r in j["rows"]] == [("EUR", 50.0), ("RSD", 400.0)]
    assert j["dimensions"] == [] and j["truncated_currencies"] == []


def test_pivot_query_group_sort_and_filter(client):
    body = {"dimensions": ["isin"], "measures": ["turnover"], "sort_by": "turnover", "order": "desc"}
    j = client.post("/api/query/pivot", json=body).json()
    assert [(r["currency"], r["dims"][0], r["values"][0]) for r in j["rows"]] == [
        ("EUR", "RSMFRSD0000C", 50.0), ("RSD", "RSMFRSD0000B", 300.0), ("RSD", "RSMFRSD0000A", 100.0)]
    j = client.post("/api/query/pivot", json={**body, "filter": {"currency": ["RSD"]}}).json()
    assert {r["currency"] for r in j["rows"]} == {"RSD"}


@pytest.mark.parametrize("body", [
    {"dimensions": ["nope"]},
    {"dimensions": ["isin", "date_day", "coupon"]},
    {"measures": []},
    {"measures": ["turnover", "turnover"]},
    {"sort_by": "quantity"},
    {"limit": 0},
    {"limit": 5001},
    {"filter": {"currency": ["rsd"]}},
    {"unknown": 1},
])
def test_pivot_query_rejects_invalid_requests(client, body):
    assert client.post("/api/query/pivot", json=body).status_code == 422


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_pivot_export_content(client, tmp_path, fmt):
    body = {"dimensions": ["isin"], "measures": ["turnover", "weighted_yield"],
            "sort_by": "turnover", "order": "desc", "format": fmt}
    j = client.post("/api/export/pivot", json=body).json()
    assert j["rows"] == 3
    p = Path(j["path"])
    assert p.exists() and p.suffix == f".{fmt}" and p.parent == tmp_path / "exports"
    header = ["Valuta", "Hartija (ISIN)", "Promet", "Ponderisani prinos"]
    if fmt == "csv":
        lines = p.read_text(encoding="utf-8-sig").splitlines()
        assert lines[0].split(",") == header and len(lines) == 4
    else:
        ws = load_workbook(p).active
        assert [c.value for c in ws[1]] == header and ws.max_row == 4
